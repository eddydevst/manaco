import csv
import io

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse

from .. import audit
from ..db import get_conn
from ..security import UsuarioActual, requiere
from ..services import importador as svc
from ..services.inventario import asegurar_fila, registrar_movimiento

router = APIRouter(prefix="/api/importador", tags=["Importador de productos"])
USAR = requiere("importador.usar")
MAX_BYTES = 5 * 1024 * 1024


async def _leer(archivo: UploadFile) -> bytes:
    contenido = await archivo.read()
    if len(contenido) > MAX_BYTES:
        raise HTTPException(413, "El archivo supera los 5 MB.")
    return contenido


@router.get("/plantilla")
def plantilla(_: UsuarioActual = Depends(USAR)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append([c[1] for c in svc.COLUMNAS])
    ws.append(["MNC-CAS-901-NG-40", "Casual", "Manaco", "Mocasín náutico", "40", "Negro",
               "Cuero con suela antideslizante", 260, 150, "PRV01", "SCZ01", 10])
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="10202B")
    for col, ancho in zip("ABCDEFGHIJKL", (22, 16, 12, 24, 8, 10, 34, 15, 10, 11, 11, 12)):
        ws.column_dimensions[col].width = ancho
    ayuda = wb.create_sheet("Instrucciones")
    ayuda.append(["Columna", "Contenido"])
    for c in svc.COLUMNAS:
        ayuda.append([c[1], c[2]])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="plantilla_productos_manaco.xlsx"'})


@router.post("/validar")
async def validar(archivo: UploadFile = File(...), conn=Depends(get_conn), _: UsuarioActual = Depends(USAR)):
    filas = svc.validar(conn, svc.leer_matriz(await _leer(archivo), archivo.filename))
    return {
        "archivo": archivo.filename,
        "resumen": {"filas": len(filas), "nuevos": sum(f["estado"] == "NUEVO" for f in filas),
                    "actualizaciones": sum(f["estado"] == "ACTUALIZA" for f in filas),
                    "errores": sum(f["estado"] == "ERROR" for f in filas)},
        "filas": svc.publico(filas),
    }


@router.post("/confirmar")
async def confirmar(archivo: UploadFile = File(...), conn=Depends(get_conn), usuario: UsuarioActual = Depends(USAR)):
    filas = svc.validar(conn, svc.leer_matriz(await _leer(archivo), archivo.filename))
    validas = [f for f in filas if f["estado"] != "ERROR"]
    errores = [f for f in filas if f["estado"] == "ERROR"]
    if not validas:
        raise HTTPException(400, "No hay filas válidas para importar.")
    id_imp = conn.execute(
        """INSERT INTO importaciones_productos (id_usuario, archivo, estado, registros_correctos, registros_error)
           VALUES (%s, %s, %s, %s, %s) RETURNING id_importacion""",
        (usuario.id_usuario, archivo.filename[:200], "CON_OBSERVACIONES" if errores else "COMPLETA", len(validas), len(errores)),
    ).fetchone()["id_importacion"]
    ref = f"IMP-{id_imp:05d}"
    sucursales = [r["id_sucursal"] for r in conn.execute("SELECT id_sucursal FROM sucursales")]
    nuevos_prod = variantes = actualizados = pares = 0
    for f in validas:
        d = f["datos"]
        if f["estado"] == "ACTUALIZA":
            antes = conn.execute("SELECT sku, precio_venta, costo_referencia FROM variantes_producto WHERE id_variante = %s FOR UPDATE",
                                 (f["_existe"]["id_variante"],)).fetchone()
            conn.execute("UPDATE variantes_producto SET precio_venta = %s, costo_referencia = %s WHERE id_variante = %s",
                         (f["_precio"], f["_costo"], f["_existe"]["id_variante"]))
            if d.get("descripcion"):
                conn.execute("UPDATE productos SET descripcion = %s WHERE id_producto = %s", (d["descripcion"], f["_existe"]["id_producto"]))
            audit.registrar(conn, usuario.id_usuario, "ACTUALIZAR_POR_IMPORTACION", "variantes_producto", f["_existe"]["id_variante"],
                            antes=antes, despues={"sku": f["sku"], "precio_venta": f["_precio"], "costo_referencia": f["_costo"], "importacion": ref})
            actualizados += 1
            continue
        prod = conn.execute(
            "SELECT id_producto FROM productos WHERE id_categoria = %s AND lower(marca) = lower(%s) AND lower(modelo) = lower(%s)",
            (f["_cat"]["id_categoria"], d["marca"], d["modelo"]),
        ).fetchone()
        if prod:
            id_prod = prod["id_producto"]
        else:
            id_prod = conn.execute(
                "INSERT INTO productos (id_categoria, marca, modelo, descripcion) VALUES (%s, %s, %s, %s) RETURNING id_producto",
                (f["_cat"]["id_categoria"], d["marca"], d["modelo"], d.get("descripcion") or None),
            ).fetchone()["id_producto"]
            nuevos_prod += 1
        if conn.execute("SELECT 1 FROM variantes_producto WHERE id_producto = %s AND lower(talla) = lower(%s) AND lower(color) = lower(%s)",
                        (id_prod, d["talla"], d["color"])).fetchone():
            raise HTTPException(409, f"Fila {f['fila']}: ya existe {d['modelo']} talla {d['talla']} color {d['color']} con otro SKU.")
        id_var = conn.execute(
            """INSERT INTO variantes_producto (id_producto, talla, color, sku, precio_venta, costo_referencia)
               VALUES (%s, %s, %s, %s, %s, %s) RETURNING id_variante""",
            (id_prod, d["talla"], d["color"], f["sku"], f["_precio"], f["_costo"]),
        ).fetchone()["id_variante"]
        variantes += 1
        for s in sucursales:
            asegurar_fila(conn, s, id_var)
        if f["_prov"]:
            conn.execute(
                """INSERT INTO productos_proveedores (id_producto, id_proveedor, codigo_proveedor, costo_referencia)
                   VALUES (%s, %s, %s, %s) ON CONFLICT DO NOTHING""",
                (id_prod, f["_prov"]["id_proveedor"], f"{f['_prov']['codigo']}-{id_prod:03d}", f["_costo"]),
            )
        if f["_stock"] > 0:
            registrar_movimiento(conn, f["_suc"]["id_sucursal"], id_var, "ENTRADA_IMPORTACION", f["_stock"], ref, usuario.id_usuario)
            pares += f["_stock"]
    for f in filas:
        conn.execute(
            "INSERT INTO detalle_importacion (id_importacion, numero_fila, sku, estado, mensaje_error) VALUES (%s, %s, %s, %s, %s)",
            (id_imp, f["fila"], f["sku"][:40] or None, f["estado"], "; ".join(f["errores"] + f["avisos"]) or None),
        )
    resultado = {"referencia": ref, "productos_nuevos": nuevos_prod, "variantes": variantes, "actualizaciones": actualizados,
                 "errores": len(errores), "pares": pares, "id_importacion": id_imp}
    audit.registrar(conn, usuario.id_usuario, "IMPORTAR_PRODUCTOS", "importaciones_productos", id_imp,
                    despues={**resultado, "archivo": archivo.filename})
    conn.commit()
    return resultado


@router.get("/historial")
def historial(conn=Depends(get_conn), _: UsuarioActual = Depends(USAR)):
    return conn.execute(
        """SELECT i.id_importacion, i.archivo, i.fecha, i.estado, i.registros_correctos, i.registros_error, u.nombre AS usuario
             FROM importaciones_productos i JOIN usuarios u USING (id_usuario) ORDER BY i.fecha DESC LIMIT 50"""
    ).fetchall()


@router.get("/{id_importacion}/errores")
def errores_csv(id_importacion: int, conn=Depends(get_conn), _: UsuarioActual = Depends(USAR)):
    filas = conn.execute(
        "SELECT numero_fila, sku, estado, mensaje_error FROM detalle_importacion WHERE id_importacion = %s AND estado = 'ERROR' ORDER BY numero_fila",
        (id_importacion,),
    ).fetchall()
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["Fila", "SKU", "Error"])
    for f in filas:
        w.writerow([f["numero_fila"], f["sku"] or "", f["mensaje_error"] or ""])
    return StreamingResponse(iter(["﻿" + buf.getvalue()]), media_type="text/csv",
                             headers={"Content-Disposition": f'attachment; filename="errores_importacion_{id_importacion}.csv"'})

