"""Lectura y validación de planillas de productos (Excel .xlsx o CSV)."""
import csv
import io
import unicodedata

from fastapi import HTTPException

COLUMNAS = [
    ("sku", "SKU", "Identificador único de la variante"),
    ("categoria", "Categoría", "Nombre o código de una categoría existente"),
    ("marca", "Marca", "Marca comercial"),
    ("modelo", "Modelo", "Modelo del producto"),
    ("talla", "Talla", "Talla del calzado"),
    ("color", "Color", "Color"),
    ("descripcion", "Descripción", "Opcional"),
    ("precio", "Precio de venta", "Número mayor a 0"),
    ("costo", "Costo", "Costo de referencia"),
    ("proveedor", "Proveedor", "Código del proveedor, ej. PRV01 (opcional)"),
    ("sucursal", "Sucursal", "Código de sucursal, ej. SCZ01 (requerido si hay stock inicial)"),
    ("stock", "Stock inicial", "Entero; solo se aplica a SKU nuevos"),
]
OBLIGATORIAS = ["sku", "categoria", "marca", "modelo", "talla", "color", "precio", "costo"]
ALIAS = {"sku": "sku", "categoria": "categoria", "marca": "marca", "modelo": "modelo", "talla": "talla", "color": "color",
         "descripcion": "descripcion", "precio de venta": "precio", "precio": "precio", "costo": "costo",
         "costo de referencia": "costo", "proveedor": "proveedor", "sucursal": "sucursal", "stock inicial": "stock", "stock": "stock"}


def norm(s) -> str:
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower().strip()
    return " ".join(s.replace("_", " ").split())


def leer_matriz(contenido: bytes, nombre: str) -> list[list[str]]:
    nombre = (nombre or "").lower()
    if nombre.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        filas = [["" if c is None else (str(int(c)) if isinstance(c, float) and c.is_integer() else str(c)) for c in fila]
                 for fila in ws.iter_rows(values_only=True)]
    elif nombre.endswith(".csv") or nombre.endswith(".txt"):
        texto = contenido.decode("utf-8-sig", errors="replace")
        primera = texto.splitlines()[0] if texto else ""
        sep = ";" if primera.count(";") >= primera.count(",") else ","
        filas = list(csv.reader(io.StringIO(texto), delimiter=sep))
    else:
        raise HTTPException(400, "Formato no soportado. Usa .xlsx o .csv.")
    return [f for f in filas if any(str(x).strip() for x in f)]


def _num(v):
    try:
        return float(str(v).replace(" ", "").replace(",", "."))
    except (TypeError, ValueError):
        return None


def validar(conn, matriz: list[list[str]]) -> list[dict]:
    if len(matriz) < 2:
        raise HTTPException(400, "El archivo no tiene filas de datos.")
    cab = [ALIAS.get(norm(h)) for h in matriz[0]]
    faltan = [k for k in OBLIGATORIAS if k not in cab]
    if faltan:
        etiquetas = {c[0]: c[1] for c in COLUMNAS}
        raise HTTPException(400, "Faltan columnas en la planilla: " + ", ".join(etiquetas[k] for k in faltan) + ". Usa la plantilla modelo.")
    categorias = {norm(r["nombre"]): r for r in conn.execute("SELECT id_categoria, codigo, nombre FROM categorias")}
    categorias.update({norm(r["codigo"]): r for r in conn.execute("SELECT id_categoria, codigo, nombre FROM categorias")})
    proveedores = {norm(r["codigo"]): r for r in conn.execute("SELECT id_proveedor, codigo FROM proveedores")}
    sucursales = {norm(r["codigo"]): r for r in conn.execute("SELECT id_sucursal, codigo FROM sucursales")}
    existentes = {r["sku"].upper(): r for r in conn.execute("SELECT id_variante, sku, id_producto FROM variantes_producto")}
    vistos, salida = {}, []
    for i, fila in enumerate(matriz[1:], start=2):
        o = {k: str(fila[j]).strip() if j < len(fila) and fila[j] is not None else "" for j, k in enumerate(cab) if k}
        err, avisos = [], []
        sku = o.get("sku", "").upper()
        if not sku:
            err.append("Falta SKU")
        elif sku in vistos:
            err.append(f"SKU duplicado en el archivo (fila {vistos[sku]})")
        else:
            vistos[sku] = i
        for k, nombre in (("marca", "marca"), ("modelo", "modelo"), ("talla", "talla"), ("color", "color")):
            if not o.get(k):
                err.append(f"Falta {nombre}")
        cat = categorias.get(norm(o.get("categoria")))
        if not o.get("categoria"):
            err.append("Falta categoría")
        elif not cat:
            err.append(f"Categoría «{o['categoria']}» no existe")
        precio, costo = _num(o.get("precio")), _num(o.get("costo"))
        if not o.get("precio"):
            err.append("Falta precio de venta")
        elif precio is None or precio <= 0:
            err.append("Precio de venta inválido")
        if not o.get("costo"):
            err.append("Falta costo")
        elif costo is None or costo < 0:
            err.append("Costo inválido")
        elif precio and costo > precio:
            avisos.append("El costo supera al precio")
        prov = None
        if o.get("proveedor"):
            prov = proveedores.get(norm(o["proveedor"]))
            if not prov:
                err.append(f"Proveedor «{o['proveedor']}» no existe")
        suc = None
        if o.get("sucursal"):
            suc = sucursales.get(norm(o["sucursal"]))
            if not suc:
                err.append(f"Sucursal «{o['sucursal']}» no existe")
        stock = 0
        if o.get("stock"):
            s = _num(o["stock"])
            if s is None or s < 0 or int(s) != s:
                err.append("Stock inicial debe ser un entero ≥ 0")
            else:
                stock = int(s)
        existe = existentes.get(sku)
        if stock > 0 and not suc and not existe:
            err.append("Indica la sucursal para el stock inicial")
        if existe and stock > 0:
            avisos.append("SKU existente: el stock no se modifica (usa compra o ajuste)")
        salida.append({
            "fila": i, "datos": o, "sku": sku, "estado": "ERROR" if err else ("ACTUALIZA" if existe else "NUEVO"),
            "errores": err, "avisos": avisos,
            "_cat": cat, "_prov": prov, "_suc": suc, "_precio": precio, "_costo": costo, "_stock": stock, "_existe": existe,
        })
    return salida


def publico(filas: list[dict]) -> list[dict]:
    return [{k: v for k, v in f.items() if not k.startswith("_")} for f in filas]
