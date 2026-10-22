"""Verificaciones de calidad y agregación sin base de datos."""
import tempfile
import unittest
from pathlib import Path
from decimal import Decimal
from openpyxl import Workbook
from backend.app.services.comercio_exterior import dashboard, leer, indicadores

HEAD = ['País_reportante', 'Flujo_comercial', 'Año', 'Código_SA6', 'Categoría_calzado', 'Socio_comercial', 'Valor_miles_USD', 'Cantidad_pares', 'Valor_USD', 'Valor_USD_por_par']

class ComercioExteriorTest(unittest.TestCase):
    def archivo(self, filas):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / 'datos.xlsx'
        w = Workbook(); w.active.append(HEAD)
        for f in filas: w.active.append(f)
        w.save(path); w.close()
        return path

    def fila(self, socio='China', usd=100, pares=10, anio=2024, flujo='Importación'):
        return ['España', flujo, anio, 640411, 'Deportivo', socio, 999, pares, usd, 999]

    def test_fuente_real_y_reconciliacion(self):
        path = Path(__file__).resolve().parents[1] / 'data' / 'comercio_exterior.xlsx'
        d = dashboard(path)
        self.assertEqual(d['calidad']['leidas'], 240)
        self.assertEqual(d['calidad']['conflictos'], 20)
        self.assertEqual(d['resumen']['registros'], 220)
        for campo in ('importaciones', 'exportaciones', 'importaciones_pares', 'exportaciones_pares'):
            self.assertAlmostEqual(sum(r[campo] for r in d['por_categoria']), d['resumen'][campo], places=4)
            self.assertAlmostEqual(sum(r[campo] for r in d['por_anio']), d['resumen'][campo], places=4)
        self.assertEqual(d['fuente']['reportantes'], ['España'])

    def test_duplicados_y_mundo_no_duplican_totales(self):
        f = self.fila()
        d = dashboard(self.archivo([f, f, self.fila('Mundo', 1000, 100)]))
        self.assertEqual(d['resumen']['importaciones'], 100)
        self.assertEqual(d['calidad']['duplicadas'], 1)
        self.assertEqual(d['calidad']['agregadas'], 1)
        self.assertEqual(len(d['agregados']), 1)

    def test_conflictos_excluye_ambas_filas(self):
        d = dashboard(self.archivo([self.fila(), self.fila(usd=200)]))
        self.assertEqual(d['calidad']['conflictos'], 2)
        self.assertEqual(d['resumen']['registros'], 0)

    def test_valor_ponderado_usa_usd_no_columnas_derivadas(self):
        d = dashboard(self.archivo([self.fila(), self.fila('Brasil', 900, 30)]))
        self.assertEqual(d['resumen']['importaciones_por_par'], 25)

    def test_nulos_nan_negativos_y_tipos(self):
        d = dashboard(self.archivo([self.fila(usd=None), self.fila(pares=-1), self.fila(anio=2024.5), self.fila(usd='NaN'), self.fila(pares='texto')]))
        self.assertEqual(d['calidad']['invalidas'], 5)
        self.assertEqual(d['resumen']['registros'], 0)

    def test_cero_pares_no_divide(self):
        d = dashboard(self.archivo([self.fila(pares=0)]))
        self.assertIsNone(d['resumen']['importaciones_por_par'])
        self.assertEqual(d['resumen']['importaciones'], 100)

    def test_filtros_y_vacio(self):
        p = self.archivo([self.fila(), self.fila('Brasil', flujo='Exportación', anio=2020)])
        d = dashboard(p, 2020, 'Brasil', 'Deportivo', 'Exportación')
        self.assertEqual(d['resumen']['exportaciones'], 100)
        self.assertEqual(d['resumen']['importaciones'], 0)
        self.assertEqual(len(d['opciones']['anio']), 2)
        self.assertIsNone(dashboard(p, pais='No existe')['mayor'])

    def test_todas_las_hojas(self):
        p = self.archivo([self.fila()])
        from openpyxl import load_workbook
        w = load_workbook(p); s = w.create_sheet('Otro flujo'); s.append(HEAD); s.append(self.fila('Brasil', flujo='Exportación')); w.save(p); w.close()
        self.assertEqual(dashboard(p)['resumen']['registros'], 2)

if __name__ == '__main__':
    unittest.main()
